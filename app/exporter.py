"""Typed, traceable Excel output. Never execute source text as a formula."""

import csv
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .engine import number, norm
from .quality import assessment

NAVY = "182D35"
TEAL = "117B70"
PALE = "EAF5F1"
AMBER = "FFF0D5"
RED = "FDE7E3"


def typed(text, header):
    key = norm(header)
    if (
        re.search(r"\b(?:id|sku|code|reference)\b", header, re.I)
        or key
        in {
            "id",
            "sku",
            "code",
            "account",
            "reference",
            "phone",
            "zip",
            "postal",
            "invoice",
            "number",
        }
        or key.endswith(
            (
                "number",
                "reference",
                "sku",
                "code",
                "itemid",
                "invoiceid",
                "accountid",
                "transactionid",
                "invoiceno",
                "accountno",
            )
        )
    ):
        return text
    if "date" in key:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
        return text
    if re.fullmatch(r"0\d+", text):
        return text
    if len(re.sub(r"[^0-9]", "", text)) > 15:
        return text  # Excel cannot represent more than 15 significant digits.
    value = number(text)
    if value is not None:
        return value / 100 if "%" in text else value
    return text


def source_number_format(text):
    """Display the source decimal precision rather than silently rounding it."""
    token = re.sub(r"[^0-9.,]", "", text)
    precision = 0
    if "." in token and "," in token:
        precision = len(token) - max(token.rfind("."), token.rfind(",")) - 1
    elif "." in token:
        precision = len(token.rsplit(".", 1)[1])
    elif "," in token and not re.fullmatch(
        r"\d{1,3}(?:,\d{3})+|\d{1,2}(?:,\d{2})*,\d{3}", token
    ):
        precision = len(token.rsplit(",", 1)[1])
    if "%" in text:
        return "0." + "0" * max(2, min(precision, 12)) + "%"
    return "#,##0" + ("." + "0" * min(precision, 12) if precision else "")


def safe_cell(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def style(ws, table_name=None):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Aptos", size=11, color=NAVY)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7F5")
        ws.row_dimensions[row[0].row].height = 30
    for col in ws.columns:
        length = max(len(str(c.value or "")) for c in list(col)[:100])
        ws.column_dimensions[col[0].column_letter].width = min(65, max(16, length + 3))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    if table_name and ws.max_row > 1:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(table)


def xlsx(result):
    wb = Workbook()
    summary = wb.active
    summary.title = "Overview"
    summary.append(["DamnLens AI Agent extraction report", "Value"])
    quality = assessment(result)
    for item in [
        ("Document", result["filename"]),
        ("Pages", result["page_count"]),
        ("Tables", len(result["tables"])),
        ("Extracted rows", result["row_count"]),
        ("Repeated headers removed", result["headers_removed"]),
        ("Validation findings", len(result["issues"])),
        ("Processing seconds", result["elapsed_seconds"]),
        ("Source SHA-256", result["sha256"]),
        ("Engine", result["engine"]),
        ("Workflow", result.get("workflow", {}).get("mode", "local")),
        ("Gemini model", result.get("workflow", {}).get("model", "Not used")),
        ("Gemini request attempts", result.get("gemini", {}).get("requests_used", 0)),
        (
            "Gemini reported total tokens",
            result.get("gemini", {})
            .get("reported_tokens", {})
            .get("totalTokenCount", 0),
        ),
        (
            "Gemini interpretation",
            "Independent unverified proposals; original data retained. Request cap is not a currency cap."
            if result.get("gemini")
            else "No external AI API requests.",
        ),
        ("Review status", quality["export_label"]),
        (
            "Pages with extracted tables",
            f"{quality['pages_extracted']} / {quality['pages_total']}",
        ),
        ("Validation errors", quality["errors"]),
        ("Validation warnings", quality["warnings"]),
        ("Rows with source evidence", quality["rows_with_evidence"]),
        ("Review note", result.get("review", {}).get("note", "")),
        ("Review completed (UTC)", result.get("review", {}).get("at", "")),
        (
            "Interpretation",
            "OCR recognition scores are not measured extraction accuracy. Review all findings.",
        ),
        (
            "Coverage",
            "Pages without extracted tables are explicitly flagged. No guarantee for arbitrary layouts.",
        ),
    ]:
        summary.append(item)
    style(summary)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 85
    if result.get("gemini"):
        candidates = wb.create_sheet("Gemini candidates")
        candidates.append(
            [
                "Page",
                "Table on page",
                "Row (0 = header)",
                "Column",
                "Header",
                "Unverified Gemini reading",
            ]
        )
        for page in result["gemini"].get("pages", []):
            for table_index, table in enumerate(page.get("tables", []), 1):
                for row_index, values in enumerate([table["headers"], *table["rows"]]):
                    for column, value in enumerate(values, 1):
                        candidates.append(
                            [
                                page["page"],
                                table_index,
                                row_index,
                                column,
                                table["headers"][column - 1],
                                None,
                            ]
                        )
                        safe_cell(
                            candidates.cell(candidates.max_row, 5),
                            table["headers"][column - 1],
                        )
                        safe_cell(candidates.cell(candidates.max_row, 6), value)
        style(candidates)
        candidates.column_dimensions["F"].width = 60
    for t in result["tables"]:
        ws = wb.create_sheet(f"Table {t['id']}")
        ws.append(
            t["headers"]
            + ["Source page", "Source row", "Extraction method", "OCR score"]
        )
        for r in t["rows"]:
            row = ws.max_row + 1
            for i, (h, value) in enumerate(zip(t["headers"], r["cells"]), 1):
                cell = safe_cell(ws.cell(row, i), typed(value, h))
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, (float, int)):
                    cell.number_format = source_number_format(value)
            for i, v in enumerate(
                [
                    r["page"],
                    r["id"],
                    r["method"],
                    r["confidence"] if r["method"].startswith("OCR") else None,
                ],
                len(t["headers"]) + 1,
            ):
                safe_cell(ws.cell(row, i), v)
        style(ws, f"Extracted{t['id']}")
        for issue in result["issues"]:
            if issue["table"] == t["id"] and issue["row"]:
                cols = (
                    [issue["column"]]
                    if issue["column"]
                    else range(1, len(t["headers"]) + 1)
                )
                for col in cols:
                    ws.cell(issue["row"] + 1, col).fill = PatternFill(
                        "solid", fgColor=RED if issue["severity"] == "error" else AMBER
                    )
    checks = wb.create_sheet("Validation")
    checks.append(["Severity", "Table", "Row", "Column", "Check", "Finding"])
    for issue in result["issues"]:
        checks.append(
            [
                issue.get(k)
                for k in ["severity", "table", "row", "column", "kind", "message"]
            ]
        )
    if not result["issues"]:
        checks.append(
            [
                "info",
                None,
                None,
                None,
                "complete",
                "No issues found by the implemented checks. Manual review remains necessary.",
            ]
        )
    style(checks)
    checks.column_dimensions["F"].width = 90
    evidence = wb.create_sheet("Evidence")
    evidence.append(
        [
            "Table",
            "Row",
            "Page",
            "Bounding box (PDF points)",
            "Method",
            "Recognition score",
        ]
    )
    for t in result["tables"]:
        for r in t["rows"]:
            for part in r.get("sources", [r]):
                evidence.append(
                    [
                        t["id"],
                        r["id"],
                        part["page"],
                        str(part["bbox"]),
                        part["method"],
                        part["confidence"]
                        if part["method"].startswith("OCR")
                        else None,
                    ]
                )
    style(evidence)
    edits = wb.create_sheet("Corrections")
    edits.append(["Table", "Row", "Column", "Original", "Edited", "Time (UTC)"])
    for e in result["edits"]:
        edits.append(
            [e.get(k) for k in ["table", "row", "column", "before", "after", "at"]]
        )
    style(edits)
    coverage = wb.create_sheet("Page coverage")
    coverage.append(
        [
            "Page",
            "Extraction",
            "Table fragments",
            "Data rows (page estimate)",
            "Seconds",
            "Status",
            "Attempts",
        ]
    )
    for p in result["pages"]:
        coverage.append(
            [
                p["page"],
                p["method"],
                p["tables"],
                p.get("rows"),
                p.get("seconds"),
                p["status"],
                "; ".join(p["attempts"]),
            ]
        )
    style(coverage)
    log = wb.create_sheet("Run log")
    log.append(["Time (UTC)", "Stage", "Action", "Page"])
    for event in result.get("events", []):
        log.append([event["at"], event["stage"], event["message"], event.get("page")])
    style(log)
    log.column_dimensions["C"].width = 70
    structural = any(
        t.get("header_paths")
        or t.get("header_ocr_disagreements")
        or any(
            r.get("merged_cells") or r.get("inherited_cells") or r.get("sources")
            for r in t["rows"]
        )
        for t in result["tables"]
    )
    if structural:
        import json

        structure = wb.create_sheet("Structure")
        structure.append(
            ["Table", "Structure", "Row", "Column", "Source page", "Details"]
        )
        for t in result["tables"]:
            for item in t.get("header_ocr_disagreements", []):
                structure.append(
                    [
                        t["id"],
                        "Header OCR disagreement",
                        None,
                        item["column"],
                        item["page"],
                        json.dumps(item, ensure_ascii=False),
                    ]
                )
            for col, path in enumerate(t.get("header_paths", []), 1):
                structure.append(
                    [t["id"], "Header hierarchy", None, col, None, " → ".join(path)]
                )
            for r in t["rows"]:
                for kind in (
                    "merged_cells",
                    "inherited_cells",
                    "sources",
                    "ocr_disagreements",
                    "ocr_retries",
                ):
                    for item in r.get(kind, []):
                        structure.append(
                            [
                                t["id"],
                                kind,
                                r["id"],
                                item.get("column"),
                                item.get("page", item.get("from_page", r["page"])),
                                json.dumps(item, ensure_ascii=False),
                            ]
                        )
        style(structure)
        structure.column_dimensions["F"].width = 90
    # Harden every sheet, including names, validation text and correction history.
    for ws in wb:
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def csv_bytes(table):
    out = io.StringIO(newline="")
    writer = csv.writer(out)

    def safe(s):
        return "'" + s if s.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else s

    writer.writerow([safe(s) for s in table["headers"]])
    writer.writerows([[safe(s) for s in r["cells"]] for r in table["rows"]])
    return out.getvalue().encode("utf-8-sig")
