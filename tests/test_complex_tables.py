import io
import json

import pymupdf as fitz
import pytest
from openpyxl import load_workbook

from app.engine import extract, record, stitch, validate
from app.exporter import xlsx
from scripts.complex_cases import additional_cases, cases

CASES = list(cases()) + list(additional_cases())


@pytest.mark.parametrize("name,data,expected", CASES, ids=[c[0] for c in CASES])
def test_complex_source_tables(tmp_path, name, data, expected):
    path = tmp_path / f"{name}.pdf"
    path.write_bytes(data)
    result = extract(path)
    actual = [
        {"headers": t["headers"], "rows": [r["cells"] for r in t["rows"]]}
        for t in result["tables"]
    ]
    assert actual == expected["tables"]
    if expected.get("header_context"):
        assert result["tables"][0].get("header_paths")
    if expected.get("merged_cells"):
        assert result["tables"][0]["rows"][1]["inherited_cells"]
        assert "merged_cell_review" in {i["kind"] for i in result["issues"]}
    if expected.get("split_row"):
        row = result["tables"][0]["rows"][2]
        assert [p["page"] for p in row["sources"]] == [1, 2]
        assert "row_continuation" in {i["kind"] for i in result["issues"]}
        book = load_workbook(io.BytesIO(xlsx(result)))
        assert book["Evidence"].max_row == 8  # six logical rows, seven source pieces
        assert "Structure" in book.sheetnames
        parts = [
            json.loads(values[5])
            for values in book["Structure"].iter_rows(min_row=2, values_only=True)
            if values[1] == "sources" and values[2] == row["id"]
        ]
        assert parts == row["sources"]
    # Every preserved source rectangle must belong to the corresponding page.
    with fitz.open(path) as document:
        for table in result["tables"]:
            for row in table["rows"]:
                for part in row.get("sources", [row]):
                    rect = fitz.Rect(part["bbox"])
                    page = document[part["page"] - 1]
                    assert rect.width > 0 and rect.height > 0
                    assert -2 <= rect.x0 < rect.x1 <= page.rect.width + 2
                    assert -2 <= rect.y0 < rect.y1 <= page.rect.height + 2


def fragment(page, rows):
    return {
        "rows": [
            record(r, page, [20, i * 20, 700, i * 20 + 20], "Native PDF")
            for i, r in enumerate(rows)
        ],
        "bbox": [20, 0, 700, 200],
        "width": 800,
    }


def test_extra_columns_are_preserved_instead_of_truncated():
    tables = stitch(
        [
            fragment(
                1,
                [["Item ID", "Amount"], ["A", "10", "extra source value"], ["B", "20"]],
            )
        ]
    )
    assert tables[0]["rows"][0]["cells"] == ["A", "10", "extra source value"]
    assert len(tables[0]["headers"]) == 3
    assert {"alignment_review", "header_review"} <= {
        i["kind"] for i in validate(tables)
    }


def test_numeric_blank_row_is_never_completed_from_arithmetic():
    tables = stitch(
        [
            fragment(
                1,
                [["Item ID", "Quantity", "Unit price", "Net"], ["A", "2", "10", "20"]],
            ),
            fragment(
                2, [["Item ID", "Quantity", "Unit price", "Net"], ["", "3", "10", ""]]
            ),
        ]
    )
    assert len(tables[0]["rows"]) == 2
    assert tables[0]["rows"][1]["cells"] == ["", "3", "10", ""]


def test_single_uncertain_ocr_cell_is_flagged_even_with_high_row_mean():
    tables = stitch([fragment(1, [["Item ID", "Amount"], ["A", "10"]])])
    row = tables[0]["rows"][0]
    row["method"] = "OCR · grid"
    row["confidence"] = 0.99
    row["cell_confidence"] = [0.99, 0.7]
    assert "low_ocr" in {i["kind"] for i in validate(tables)}


def test_separately_drawn_identifier_glyphs_keep_geometric_order(tmp_path):
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(400, 400))
    for x in (25, 175, 325):
        c.line(x, 200, x, 300)
    for y in (200, 250, 300):
        c.line(25, y, 325, y)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(30, 280, "Item ID")
    c.drawString(180, 280, "Amount")
    c.setFont("Helvetica", 12)
    c.drawString(30, 230, "AB")
    c.drawString(52, 230, "01")
    c.drawString(46, 230, "_")
    c.drawString(180, 230, "45.00")
    c.save()
    path = tmp_path / "glyphs.pdf"
    path.write_bytes(stream.getvalue())
    assert extract(path)["tables"][0]["rows"][0]["cells"] == ["AB_01", "45.00"]


@pytest.mark.parametrize("borderless", [False, True])
def test_business_rows_with_header_words_are_never_dropped(tmp_path, borderless):
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(600, 500))
    headers = ["Name", "Department", "Status"]
    rows = [
        ["Sales manager", "Revenue office", "Open"],
        ["Account manager", "Product team", "Active"],
    ]
    if not borderless:
        for x in (25, 195, 365, 535):
            c.line(x, 250, x, 350)
        for y in (250, 285, 320, 350):
            c.line(25, y, 535, y)
    for ri, row in enumerate([headers] + rows):
        c.setFont("Helvetica-Bold" if ri == 0 else "Helvetica", 11)
        for ci, value in enumerate(row):
            c.drawString(30 + ci * 170, 330 - ri * 35, value)
    c.save()
    path = tmp_path / "words.pdf"
    path.write_bytes(stream.getvalue())
    result = extract(path)
    assert result["tables"][0]["headers"] == headers
    assert [r["cells"] for r in result["tables"][0]["rows"]] == rows
