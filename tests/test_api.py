import importlib
import io
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("LEDGERLENS_DATA", str(tmp_path / "jobs"))
    import app.main as module

    module = importlib.reload(module)
    with TestClient(module.app) as client:
        yield client
    module.POOL.shutdown(wait=True)


def completed(client):
    response = client.post("/api/samples/01_invoice_register")
    assert response.status_code == 202
    job = response.json()["id"]
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        state = client.get("/api/jobs/" + job).json()
        if state["status"] in {"complete", "failed"}:
            break
        time.sleep(0.05)
    assert state["status"] == "complete"
    return job


def test_health_and_frontend(client):
    assert client.get("/api/health").json()["max_pages"] == 50
    assert "Document extraction" in client.get("/").text
    assert client.get("/styles.css").status_code == 200
    assert client.get("/api/health").headers["X-Content-Type-Options"] == "nosniff"


def test_job_export_correct_delete_and_history(client):
    job = completed(client)
    assert len(client.get("/api/jobs").json()) == 1
    r = client.get(f"/api/jobs/{job}/result").json()
    assert r["row_count"] == 31
    patch = client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "999.00"},
    )
    assert patch.status_code == 200
    r = patch.json()
    assert r["edits"][0]["before"] == "44.00" and r["edits"][0]["after"] == "999.00"
    assert {"arithmetic", "total_mismatch"} <= {i["kind"] for i in r["issues"]}
    assert client.get(f"/api/jobs/{job}/export/xlsx").content[:2] == b"PK"
    assert client.get(f"/api/jobs/{job}/export/json").json()["edits"]
    assert client.get(f"/api/jobs/{job}/export/csv").status_code == 200
    assert client.get(f"/api/jobs/{job}/page/3").headers["content-type"] == "image/png"
    assert client.get(f"/api/jobs/{job}/page/4").status_code == 404
    assert client.delete("/api/jobs/" + job).status_code == 204
    assert client.get("/api/jobs/" + job).status_code == 404


@pytest.mark.parametrize(
    "name,data,status",
    [
        ("bad.txt", b"hello", 415),
        ("bad.pdf", b"%PDF-invalid", 422),
        ("large.pdf", b"0" * (40 * 1024 * 1024 + 1), 413),
    ],
    ids=["unsupported-type", "corrupt-pdf", "oversize-upload"],
)
def test_upload_rejections(client, name, data, status):
    assert client.post("/api/jobs", files={"file": (name, data)}).status_code == status


def test_path_and_origin_security(client):
    assert client.get("/api/jobs/not-a-uuid").status_code == 404
    assert (
        client.post(
            "/api/samples/01_invoice_register", headers={"origin": "https://evil.test"}
        ).status_code
        == 403
    )
    assert client.get("/api/health", headers={"host": "evil.test"}).status_code == 403
    assert client.post("/api/samples/no-such-file").status_code == 404


def test_invalid_cell_coordinates_rejected(client):
    job = completed(client)
    assert (
        client.patch(
            f"/api/jobs/{job}/cells",
            json={"table": 1, "row": 0, "column": 1, "value": "x"},
        ).status_code
        == 422
    )
    assert (
        client.patch(
            f"/api/jobs/{job}/cells",
            json={"table": 99, "row": 1, "column": 1, "value": "x"},
        ).status_code
        == 404
    )


def test_restart_marks_interrupted(client):
    import app.main as module

    job = "00000000-0000-0000-0000-000000000000"
    p = module.JOBS / job
    p.mkdir()
    module.save(
        p / "status.json", {"id": job, "status": "processing", "created": "2026-10-15"}
    )
    importlib.reload(module)
    assert module.load(p / "status.json")["status"] == "interrupted"


def test_review_is_explicit_and_corrections_revoke_it(client):
    job = completed(client)
    assert (
        client.get(f"/api/jobs/{job}/result").json()["quality"]["review_status"]
        == "pending"
    )
    assert client.post(f"/api/jobs/{job}/review", json={}).status_code == 422
    r = client.post(
        f"/api/jobs/{job}/review",
        json={"source_checked": True, "note": "Compared source and columns."},
    ).json()
    assert r["quality"]["review_status"] == "reviewed"
    assert r["review"]["fingerprint"] and r["review"]["accepted_findings"] == 0
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job}/export/xlsx").content))
    overview = dict(wb["Overview"].iter_rows(min_row=2, values_only=True))
    assert (
        overview["Review status"] == "Reviewed"
        and overview["Review note"] == "Compared source and columns."
    )
    same = client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "44.00"},
    ).json()
    assert same["quality"]["review_status"] == "reviewed" and same["edits"] == []
    r = client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "999"},
    ).json()
    assert r["quality"]["review_status"] == "pending"
    assert r["review_history"][0]["note"] == "Compared source and columns."
    assert client.get("/api/jobs").json()[0]["review_status"] == "pending"
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job}/export/xlsx").content))
    assert (
        dict(wb["Overview"].iter_rows(min_row=2, values_only=True))["Review status"]
        == "Draft — review pending"
    )


def test_outstanding_findings_require_acknowledgement_and_note(client):
    job = completed(client)
    client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "999"},
    )
    route = f"/api/jobs/{job}/review"
    assert client.post(route, json={"source_checked": True}).status_code == 409
    assert (
        client.post(
            route, json={"source_checked": True, "accept_findings": True}
        ).status_code
        == 409
    )
    r = client.post(
        route,
        json={
            "source_checked": True,
            "accept_findings": True,
            "note": "Synthetic source mismatch accepted for workflow test.",
        },
    ).json()
    assert (
        r["quality"]["review_status"] == "reviewed"
        and r["review"]["accepted_findings"] > 0
    )
    assert client.get("/api/review").json()[0]["review_status"] == "reviewed"


def test_retry_archives_corrections_and_blocks_stale_exports(client, monkeypatch):
    import app.main as module

    job = completed(client)
    client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "999"},
    )
    monkeypatch.setattr(module.POOL, "submit", lambda *args: None)
    state = client.post(f"/api/jobs/{job}/retry?force_ocr=true").json()
    assert state["status"] == "queued" and state["attempt"] == 2
    assert client.get(f"/api/jobs/{job}/result").status_code == 409
    assert client.get(f"/api/jobs/{job}/export/xlsx").status_code == 409
    versions = client.get(f"/api/jobs/{job}/versions").json()
    assert len(versions) == 1 and versions[0]["edits"] == 1
    snapshot = client.get(
        f"/api/jobs/{job}/versions/{versions[0]['id']}/export/json"
    ).json()
    assert snapshot["tables"][0]["rows"][0]["cells"][7] == "999"
    assert (
        client.get(f"/api/jobs/{job}/versions/{versions[0]['id']}/export/xlsx").content[
            :2
        ]
        == b"PK"
    )


def test_inbox_includes_document_and_source_page(client):
    job = completed(client)
    client.patch(
        f"/api/jobs/{job}/cells",
        json={"table": 1, "row": 1, "column": 8, "value": "999"},
    )
    inbox = client.get("/api/review").json()
    assert any(i["job"] == job and i["page"] == 1 and i["row"] == 1 for i in inbox)


def test_progress_write_recovers_from_transient_windows_file_locks(
    client, tmp_path, monkeypatch
):
    import app.main as module

    target = tmp_path / "state.json"
    target.write_text('{"status":"old"}')
    original = Path.replace
    attempts = []

    def temporarily_locked(path, destination):
        attempts.append(path)
        if len(attempts) < 3:
            raise PermissionError("simulated Windows sharing lock")
        return original(path, destination)

    monkeypatch.setattr(Path, "replace", temporarily_locked)
    module.save(target, {"status": "complete"})
    assert module.load(target) == {"status": "complete"} and len(attempts) == 3
    assert list(tmp_path.glob("*.tmp")) == []
