"""Cloud contract/recovery tests use mocked responses, never a live API key."""

from copy import deepcopy
import io
import json
from pathlib import Path
import urllib.error

import pymupdf as fitz
import pytest
from openpyxl import load_workbook

from app import gemini, workflow
from app.exporter import xlsx
from app.quality import fingerprint

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def isolate_key(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_MODEL", raising=False)


def baseline():
    r = json.loads(
        (ROOT / "reports/01_invoice_register.json").read_text(encoding="utf-8")
    )
    r.pop("review", None)
    return r


def response(tables, finish="STOP", **extra):
    return {
        "candidates": [
            {
                "finishReason": finish,
                "content": {"parts": [{"text": json.dumps({"tables": tables})}]},
            }
        ],
        "usageMetadata": {
            "promptTokenCount": 100,
            "candidatesTokenCount": 80,
            "totalTokenCount": 180,
        },
        **extra,
    }


def candidate(r, page):
    return [
        {
            "headers": t["headers"],
            "rows": [row["cells"][:] for row in t["rows"] if row["page"] == page],
        }
        for t in r["tables"]
        if any(row["page"] == page for row in t["rows"])
    ]


def config(monkeypatch, limit=3):
    monkeypatch.setenv("GEMINI_API_KEY", "dummy-test-key-not-real")
    return workflow.options("gemini", True, "all", limit)


def save(path, value):
    path.write_text(json.dumps(value), encoding="utf-8")


def test_local_never_calls_cloud(monkeypatch):
    def forbidden(*args):
        pytest.fail("Local workflow attempted a cloud call")

    monkeypatch.setattr(gemini, "request_page", forbidden)
    monkeypatch.setattr(workflow, "extract", lambda *a, **k: baseline())
    r = workflow.run(ROOT / "samples/01_invoice_register.pdf")
    assert r["workflow"] == {"mode": "local"} and "gemini" not in r


@pytest.mark.parametrize(
    "mode,consent,scope,limit",
    [
        ("unknown", True, "all", 3),
        ("gemini", False, "all", 3),
        ("gemini", True, "unknown", 3),
        ("gemini", True, "all", 0),
        ("gemini", True, "all", 51),
        ("gemini", True, "all", True),
    ],
)
def test_reject_invalid_authorization_and_limits(
    monkeypatch, mode, consent, scope, limit
):
    monkeypatch.setenv("GEMINI_API_KEY", "dummy")
    with pytest.raises(ValueError):
        workflow.options(mode, consent, scope, limit)


def test_missing_key_and_invalid_model(monkeypatch):
    with pytest.raises(ValueError, match="Configure"):
        workflow.options("gemini", True)
    monkeypatch.setenv("GEMINI_API_KEY", "dummy")
    monkeypatch.setenv("GEMINI_MODEL", "https://example.invalid/steal")
    assert not gemini.configuration()["model_valid"]
    with pytest.raises(ValueError):
        workflow.options("gemini", True)


def test_environment_is_data_not_code(tmp_path, monkeypatch):
    (tmp_path / ".env").write_text(
        'GEMINI_API_KEY="dummy"\nGEMINI_MODEL=gemini-test\nPATH=bad\nHOME=bad\n',
        encoding="utf-8",
    )
    gemini.load_environment(tmp_path)
    assert gemini.configuration()["configured"]
    monkeypatch.setenv("GEMINI_API_KEY", "environment-precedence")
    gemini.load_environment(tmp_path)
    import os

    assert os.environ["GEMINI_API_KEY"] == "environment-precedence"
    assert os.environ.get("PATH") != "bad" and os.environ.get("HOME") != "bad"


@pytest.mark.parametrize(
    "bad",
    [
        response([], "MAX_TOKENS"),
        response([], "SAFETY"),
        {},
        response([{"headers": ["A", "B"], "rows": [["1"]]}]),
        response([{"headers": ["A", "B"], "rows": [[1, "2"]]}]),
        response([{"headers": ["A", "B"], "rows": [["x" * 2001, "2"]]}]),
        response([{"headers": ["A", "B"], "rows": [], "instructions": "run code"}]),
        response([], usageMetadata={"totalTokenCount": -1}),
    ],
)
def test_untrusted_and_truncated_output_is_rejected(bad):
    with pytest.raises(gemini.GeminiError):
        gemini.parse_response(bad)


def test_native_local_difficult_page_selection():
    r = baseline()
    assert gemini.selected_pages(r, "flagged") == []
    r["pages"][1]["method"] = "ocr"
    assert gemini.selected_pages(r, "flagged") == [2]
    r["issues"] = [{"table": 1, "row": 1}]
    assert gemini.selected_pages(r, "flagged") == [1, 2]


def test_single_page_request_contract_and_no_tools(monkeypatch):
    config(monkeypatch)
    captured = {}

    class Opener:
        def open(self, req, timeout):
            captured.update(
                url=req.full_url,
                headers=req.headers,
                payload=json.loads(req.data),
                timeout=timeout,
            )
            return io.BytesIO(json.dumps(response([])).encode())

    monkeypatch.setattr(gemini.urllib.request, "build_opener", lambda *a: Opener())
    gemini.request_page(b"%PDF-dummy", gemini.DEFAULT_MODEL)
    assert captured["url"].startswith("https://generativelanguage.googleapis.com/")
    assert "dummy-test-key" not in captured["url"]
    assert captured["timeout"] == 60 and "tools" not in captured["payload"]
    assert captured["payload"]["generationConfig"]["maxOutputTokens"] == 8192
    assert (
        "untrusted data" in captured["payload"]["systemInstruction"]["parts"][0]["text"]
    )
    assert (
        gemini.NoRedirect().redirect_request(
            None, None, 302, "", {}, "https://example.invalid"
        )
        is None
    )


@pytest.mark.parametrize("code", [400, 401, 403, 429, 500, 503])
def test_http_errors_never_leak_credentials(monkeypatch, code):
    config(monkeypatch)

    class Opener:
        def open(self, *a, **k):
            raise urllib.error.HTTPError(
                "https://secret.invalid/dummy-test-key",
                code,
                "secret provider message",
                {},
                None,
            )

    monkeypatch.setattr(gemini.urllib.request, "build_opener", lambda *a: Opener())
    with pytest.raises(gemini.GeminiError) as error:
        gemini.request_page(b"pdf", gemini.DEFAULT_MODEL)
    assert str(code) in str(error.value) and "secret" not in str(error.value)


def test_suggestions_preserve_originals_and_duplicate_ambiguity():
    r = baseline()
    c = candidate(r, 1)
    original = deepcopy(r)
    c[0]["rows"][0][7] = "999.00"
    proposals, structural = gemini.compare_page(r, 1, c)
    assert len(proposals) == 1 and not structural
    assert proposals[0]["before"] == "44.00" and proposals[0]["suggested"] == "999.00"
    assert r == original
    r["tables"][0]["rows"].insert(1, deepcopy(r["tables"][0]["rows"][0]))
    c = candidate(r, 1)
    c[0]["rows"][0][7] = "999.00"
    proposals, structural = gemini.compare_page(r, 1, c)
    assert not proposals and structural  # Cannot map repeated anchors safely.


def test_reordered_or_missing_rows_are_not_mapped():
    r = baseline()
    c = candidate(r, 1)
    c[0]["rows"].reverse()
    assert gemini.compare_page(r, 1, c)[1]
    c[0]["rows"].pop()
    assert gemini.compare_page(r, 1, c) == ([], True)


def test_budget_reservation_resume_usage_and_exports(tmp_path, monkeypatch):
    opt = config(monkeypatch, 1)
    path = tmp_path / "ledger.json"
    calls = []
    r = baseline()
    original = deepcopy(r["tables"])

    def fake(pdf, model, mime_type):
        assert mime_type == "application/pdf"
        assert json.loads(path.read_text())["records"][0]["status"] == "reserved"
        with fitz.open("pdf", pdf) as page:
            assert len(page) == 1
        calls.append(model)
        tables = candidate(r, 1)
        tables[0]["rows"][0][7] = "999.00"
        return response(tables)

    monkeypatch.setattr(gemini, "request_page", fake)
    monkeypatch.setattr(workflow, "extract", lambda *a, **k: deepcopy(r))
    output = workflow.run(ROOT / "samples/01_invoice_register.pdf", opt, path, save)
    assert output["tables"] == original
    assert output["gemini"]["requests_used"] == 1
    assert output["gemini"]["reported_tokens"]["totalTokenCount"] == 180
    assert [p["status"] for p in output["gemini"]["pages"]] == [
        "checked",
        "skipped",
        "skipped",
    ]
    assert {"gemini_disagreement", "gemini_coverage"} <= {
        i["kind"] for i in output["issues"]
    }
    workflow.run(ROOT / "samples/01_invoice_register.pdf", opt, path, save)
    assert len(calls) == 1  # Resume never duplicates an already charged call.
    assert "dummy-test-key" not in json.dumps(output) + path.read_text()
    workbook = load_workbook(io.BytesIO(xlsx(output)))
    assert "Gemini candidates" in workbook.sheetnames
    assert workbook["Table 1"].cell(2, 8).value == 44
    before = fingerprint(output)
    output["gemini"]["pages"][0]["proposals"][0]["suggested"] = "42"
    assert fingerprint(output) != before


def test_interrupted_reserved_request_not_resent(tmp_path, monkeypatch):
    opt = config(monkeypatch)
    r = baseline()
    path = tmp_path / "ledger.json"

    def fail(*a):
        raise RuntimeError("dummy-test-key-secret")

    monkeypatch.setattr(gemini, "request_page", fail)
    with fitz.open(ROOT / "samples/01_invoice_register.pdf") as doc:
        first = gemini.assist(r, doc, opt, lambda *a, **k: None, path, save)
        assert first["pages"][0]["status"] == "failed"
        assert "secret" not in json.dumps(first)
        ledger = json.loads(path.read_text())
        ledger["records"][0] = {"page": 1, "status": "reserved"}
        import hashlib

        ledger["checksum"] = hashlib.sha256(
            json.dumps(ledger["records"], sort_keys=True).encode()
        ).hexdigest()
        save(path, ledger)
        monkeypatch.setattr(
            gemini, "request_page", lambda *a: pytest.fail("Unknown request resent")
        )
        again = gemini.assist(r, doc, opt, lambda *a, **k: None, path, save)
        assert again["pages"][0]["status"] == "unknown" and again["requests_used"] == 1


def test_ledger_failure_and_changed_baseline_send_nothing(tmp_path, monkeypatch):
    opt = config(monkeypatch)
    r = baseline()
    path = tmp_path / "ledger.json"
    monkeypatch.setattr(
        gemini, "request_page", lambda *a: pytest.fail("Unreserved call")
    )

    def disk_failure(*a):
        raise OSError("private path")

    monkeypatch.setattr(workflow, "extract", lambda *a, **k: deepcopy(r))
    output = workflow.run(
        ROOT / "samples/01_invoice_register.pdf", opt, path, disk_failure
    )
    assert (
        output["tables"] == r["tables"] and output["gemini"]["status"] == "unavailable"
    )
    assert "private path" not in json.dumps(output)
    save(path, {"signature": "different-baseline", "records": []})
    output = workflow.run(ROOT / "samples/01_invoice_register.pdf", opt, path, save)
    assert "baseline changed" in output["gemini"]["message"]


def test_candidate_formula_is_exported_as_text(monkeypatch):
    r = baseline()
    config(monkeypatch)
    r["gemini"] = {
        "pages": [
            {
                "page": 1,
                "tables": [
                    {"headers": ["=evil", "B"], "rows": [['=HYPERLINK("bad")', "text"]]}
                ],
            }
        ]
    }
    workbook = load_workbook(io.BytesIO(xlsx(r)))
    assert workbook["Gemini candidates"]["F4"].data_type == "s"
    assert all(c.data_type != "f" for row in workbook["Gemini candidates"] for c in row)


def test_corrupt_ledger_refuses_new_charges(tmp_path, monkeypatch):
    opt = config(monkeypatch)
    path = tmp_path / "ledger.json"
    r = baseline()
    monkeypatch.setattr(gemini, "request_page", lambda *a: response([]))
    with fitz.open(ROOT / "samples/01_invoice_register.pdf") as doc:
        gemini.assist(r, doc, opt, lambda *a, **k: None, path, save)
        ledger = json.loads(path.read_text())
        ledger["records"][0]["page"] = 2
        save(path, ledger)
        monkeypatch.setattr(
            gemini,
            "request_page",
            lambda *a: pytest.fail("Corrupt ledger sent request"),
        )
        with pytest.raises(gemini.GeminiError, match="verification failed"):
            gemini.assist(r, doc, opt, lambda *a, **k: None, path, save)


def test_scanned_pages_use_bounded_png_input(monkeypatch):
    opt = config(monkeypatch, 1)
    r = baseline()
    r["pages"][0]["method"] = "ocr"

    def provider(media, model, mime_type):
        assert media.startswith(b"\x89PNG") and mime_type == "image/png"
        return response(candidate(r, 1))

    monkeypatch.setattr(gemini, "request_page", provider)
    with fitz.open(ROOT / "samples/01_invoice_register.pdf") as doc:
        output = gemini.assist(r, doc, opt, lambda *a, **k: None)
    assert output["pages"][0]["status"] == "checked"
    assert output["pages"][0]["input_mime_type"] == "image/png"
