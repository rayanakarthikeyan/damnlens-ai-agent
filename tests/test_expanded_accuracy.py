import io
import json
from pathlib import Path

import numpy as np
import pymupdf as fitz
import pytest
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle

from app import ocr_verify
from app.engine import extract, record, stitch, validate
from app.exporter import xlsx
from app.layout import reading_order


def pdf_table(data, styles=()):
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(600, 400))
    table = Table(data, colWidths=[150] * len(data[0]), rowHeights=[35] * len(data))
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.7, colors.black),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 12),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                *styles,
            ]
        )
    )
    table.wrapOn(c, 500, 350)
    table.drawOn(c, 35, 300 - 35 * len(data))
    c.save()
    return stream.getvalue()


def save_document(tmp_path, data, scanned=False):
    path = tmp_path / "source.pdf"
    if scanned:
        with fitz.open(stream=data, filetype="pdf") as source, fitz.open() as target:
            for page in source:
                image = page.get_pixmap(dpi=180, alpha=False)
                p = target.new_page(width=page.rect.width, height=page.rect.height)
                p.insert_image(p.rect, stream=image.tobytes("png"))
            target.save(path)
    else:
        path.write_bytes(data)
    return path


def test_column_reference_row_is_header_context_not_a_business_row(tmp_path):
    data = pdf_table(
        [
            ["Product", "Quantity", "Amount"],
            ["(1)", "(2)", "(3)"],
            ["Cable", "2", "10.00"],
            ["Plug", "4", "20.00"],
        ]
    )
    result = extract(save_document(tmp_path, data))
    t = result["tables"][0]
    assert t["headers"] == ["Product", "Quantity", "Amount"]
    assert [r["cells"] for r in t["rows"]] == [
        ["Cable", "2", "10.00"],
        ["Plug", "4", "20.00"],
    ]
    assert t["header_rows"][1] == ["(1)", "(2)", "(3)"]


def test_parenthesized_business_amounts_are_not_removed():
    rows = [
        record(v, 1, [10, i * 25, 300, i * 25 + 25], "Native PDF")
        for i, v in enumerate(
            [["Product", "Quantity", "Amount"], ["Return", "(2)", "(3)"]]
        )
    ]
    result = stitch([{"rows": rows, "bbox": [10, 0, 300, 50], "width": 400}])
    assert result[0]["rows"][0]["cells"] == ["Return", "(2)", "(3)"]


def test_external_header_words_are_not_clipped(tmp_path):
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(600, 400))
    c.setFont("Helvetica", 12)
    for x, value in zip(
        (59.6, 225.8, 392), ("Hello", "Goodbye", "World"), strict=False
    ):
        c.drawString(x, 330, value)
    for i, row in enumerate(
        [
            ["1", "Abc", "Foo"],
            ["2", "Bar", "Baz"],
            ["3", "Quux", "Argh"],
            ["4", "Chonk", "Bonk"],
        ]
    ):
        for x, value in zip((56.8, 223, 389.2), row, strict=False):
            c.drawString(x, 313 - i * 14, value)
    c.save()
    result = extract(save_document(tmp_path, stream.getvalue()))
    assert result["tables"][0]["headers"] == ["Hello", "Goodbye", "World"]
    assert len(result["tables"][0]["rows"]) == 4


def test_header_only_region_has_no_fabricated_rows_and_needs_review(tmp_path):
    data = pdf_table(
        [["Pending report", "", ""], ["Product", "Quantity", "Amount"]],
        [("SPAN", (0, 0), (-1, 0))],
    )
    result = extract(save_document(tmp_path, data))
    assert result["tables"] == []
    assert result["pages"][0]["status"] == "needs_review"
    assert result["quality"]["pages_extracted"] == 0
    assert any(i["kind"] == "page_review" for i in result["issues"])


@pytest.mark.parametrize("scanned", [False, True], ids=["native", "scanned"])
@pytest.mark.parametrize(
    "merged", [False, True], ids=["intentional-blanks", "explicit-numeric-span"]
)
def test_numeric_shared_values_need_explicit_geometry(tmp_path, scanned, merged):
    source = [
        ["Batch", "Item", "Quantity"],
        ["4", "Cable", "2"],
        ["", "Plug", "4"],
        ["", "Adapter", "1"],
    ]
    styles = [("SPAN", (0, 1), (0, 3))] if merged else []
    result = extract(save_document(tmp_path, pdf_table(source, styles), scanned))
    rows = result["tables"][0]["rows"]
    assert [r["cells"][0] for r in rows] == (["4"] * 3 if merged else ["4", "", ""])
    assert [r["cells"][1:] for r in rows] == [
        ["Cable", "2"],
        ["Plug", "4"],
        ["Adapter", "1"],
    ]
    if merged:
        assert all(r.get("inherited_cells") for r in rows[1:])
        assert all(
            i["bbox"][1] < rows[1]["bbox"][1] for i in rows[1]["inherited_cells"]
        )


def test_cell_reading_order_handles_baseline_jitter():
    words = [
        {"text": v, "bbox": [x, y, x + 30, y + 10]}
        for v, x, y in [
            ("Cases", 10, 7),
            ("reported", 45, 8),
            ("from", 100, 7),
            ("Village", 140, 5),
            ("Sankripali", 10, 22),
        ]
    ]
    assert [w["text"] for w in reading_order(words)] == [
        "Cases",
        "reported",
        "from",
        "Village",
        "Sankripali",
    ]


@pytest.mark.parametrize(
    "original,candidates,expected,uncertain",
    [
        ("ii. Food", [["iii. Food", 0.97], ["iii. Food", 0.98]], "iii. Food", False),
        ("SI.", [["Sl.", 0.85], ["SI.", 0.85]], "SI.", True),
        ("No.", [["NO.", 0.99], ["NO.", 0.99]], "No.", False),
    ],
)
def test_secondary_ocr_preserves_disagreement_and_only_accepts_consensus(
    monkeypatch, original, candidates, expected, uncertain
):
    replies = iter(candidates)
    monkeypatch.setattr(ocr_verify, "_model", lambda *a, **k: ([next(replies)], None))
    words = [{"text": original, "score": 0.999, "bbox": [5, 5, 90, 30]}]
    actual = ocr_verify.verify_labels(
        np.full((40, 100, 3), 255, dtype=np.uint8), words
    )[0]
    assert actual["text"] == expected
    assert bool(actual.get("disagreement")) == uncertain
    if uncertain:
        assert actual["score"] < 0.88
        assert [c["text"] for c in actual["disagreement"]["candidates"]] == [
            "Sl.",
            "SI.",
        ]
    if expected != original:
        assert actual["verification"]["original"] == original


def test_header_disagreement_survives_validation_and_excel():
    rows = [
        record(v, 1, [10, i * 25, 300, i * 25 + 25], "OCR · grid")
        for i, v in enumerate([["SI. No.", "Year"], ["1", "2024"], ["2", "2025"]])
    ]
    rows[0]["_header_hint"] = True
    rows[0]["ocr_disagreements"] = [
        {
            "column": 1,
            "original": "SI.",
            "bbox": [10, 0, 40, 25],
            "candidates": [
                {"text": "Sl.", "score": 0.85},
                {"text": "SI.", "score": 0.85},
            ],
        }
    ]
    tables = stitch([{"rows": rows, "bbox": [10, 0, 300, 75], "width": 400}])
    issues = validate(tables)
    assert any(
        i["kind"] == "ocr_disagreement" and i["page"] == 1 and i["column"] == 1
        for i in issues
    )
    # Use a complete real result envelope for export metadata.
    source = Path(__file__).resolve().parents[1] / "samples/01_invoice_register.pdf"
    result = extract(source)
    result.update(tables=tables, issues=issues, row_count=2)
    book = load_workbook(io.BytesIO(xlsx(result)))
    entries = [
        r
        for r in book["Structure"].iter_rows(min_row=2, values_only=True)
        if r[1] == "Header OCR disagreement"
    ]
    assert (
        len(entries) == 1
        and json.loads(entries[0][5])["candidates"][0]["text"] == "Sl."
    )


@pytest.mark.parametrize(
    "shared", [True, False], ids=["one-merged-source", "three-distinct-sources"]
)
def test_totals_count_shared_source_amount_once_without_deduplicating_business_rows(
    shared,
):
    rows = [
        record([name, "4"], 1, [10, i * 25, 300, i * 25 + 25], "Native PDF")
        for i, name in enumerate(["A", "B", "C"])
    ]
    if shared:
        rows[0]["merged_cells"] = [
            {
                "column": 2,
                "column_span": 1,
                "row_span": 3,
                "bbox": [150, 0, 300, 75],
                "value": "4",
            }
        ]
        for row in rows[1:]:
            row["inherited_cells"] = [
                {
                    "column": 2,
                    "from_page": 1,
                    "from_grid_row": 1,
                    "bbox": [150, 0, 300, 75],
                    "value": "4",
                }
            ]
    rows.append(
        record(["Total", "4" if shared else "12"], 1, [10, 75, 300, 100], "Native PDF")
    )
    for i, row in enumerate(rows, 1):
        row["id"] = i
    tables = [{"id": 1, "headers": ["Item", "Amount"], "rows": rows}]
    assert not any(i["kind"] == "total_mismatch" for i in validate(tables))
    if shared:
        assert any(i["kind"] == "merged_total_review" for i in validate(tables))
        rows[-1]["cells"][1] = "12"
        assert not any(i["kind"] == "total_mismatch" for i in validate(tables))
    rows[-1]["cells"][1] = "5"
    assert any(i["kind"] == "total_mismatch" for i in validate(tables))
