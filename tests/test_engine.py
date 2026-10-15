import io
import json
from copy import deepcopy
from pathlib import Path

import pymupdf as fitz
import pytest
from openpyxl import load_workbook

from app.engine import extract, number, record, stitch, validate
from app.exporter import xlsx, typed, csv_bytes
from scripts.generate_samples import invoice_fixture, render_pdf

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize(
    "text,expected",
    [
        ("1,234.50", 1234.5),
        ("1.234,50", 1234.5),
        ("₹ 1,20,000.25", 120000.25),
        ("(75.50)", -75.5),
        ("0", 0),
        ("10%", 10),
        ("EUR 42,50", 42.5),
        ("2026-10-01", None),
        ("INV-0001", None),
        ("", None),
        ("—", None),
    ],
)
def test_number_locales(text, expected):
    assert number(text) == expected


def fragment(page, rows):
    return {
        "rows": [
            record(r, page, [40, 100 + i * 20, 700, 120 + i * 20], "Native PDF")
            for i, r in enumerate(rows)
        ],
        "bbox": [40, 100, 700, 300],
        "width": 800,
    }


def test_repeated_headers_removed_but_valid_duplicates_preserved():
    h = ["Item ID", "Description", "Amount"]
    r = ["01", "Duplicate business record", "10"]
    tables = stitch([fragment(1, [h, r]), fragment(2, [h, r])])
    assert len(tables) == 1 and len(tables[0]["rows"]) == 2
    assert tables[0]["headers_removed"] == 1
    assert "duplicate_row" in {i["kind"] for i in validate(tables)}


def test_headerless_continuation_keeps_first_data_row():
    h = ["Item ID", "Description", "Amount"]
    t = stitch(
        [
            fragment(1, [h, ["01", "First", "10"]]),
            fragment(2, [["02", "Second", "20"], ["03", "Third", "30"]]),
        ]
    )
    assert len(t) == 1 and [r["cells"][0] for r in t[0]["rows"]] == ["01", "02", "03"]


def test_nonadjacent_and_different_schemas_not_joined():
    h = ["Item ID", "Description", "Amount"]
    assert (
        len(
            stitch(
                [
                    fragment(1, [h, ["01", "One", "1"]]),
                    fragment(3, [h, ["02", "Two", "2"]]),
                ]
            )
        )
        == 2
    )
    assert (
        len(
            stitch(
                [
                    fragment(1, [h, ["01", "One", "1"]]),
                    fragment(
                        2, [["Date", "Debit", "Balance"], ["2026-01-01", "2", "3"]]
                    ),
                ]
            )
        )
        == 2
    )


def test_digital_fixture_exact_cells_and_provenance():
    expected = json.loads((ROOT / "samples" / "ground_truth.json").read_text())[
        "01_invoice_register.pdf"
    ]
    r = extract(ROOT / "samples" / "01_invoice_register.pdf")
    assert len(r["tables"]) == 1 and r["headers_removed"] == 2
    assert r["tables"][0]["headers"] == expected["headers"]
    assert [row["cells"] for row in r["tables"][0]["rows"]] == expected["rows"]
    assert r["issues"] == []
    assert all(
        len(row["bbox"]) == 4 and 1 <= row["page"] <= 3
        for row in r["tables"][0]["rows"]
    )


def test_bank_errors_are_preserved_and_detected():
    r = extract(ROOT / "samples" / "03_bank_statement.pdf")
    kinds = {i["kind"] for i in r["issues"]}
    assert {"total_mismatch", "balance_mismatch", "missing_value"} <= kinds
    assert r["tables"][0]["rows"][8]["cells"][7] == "6050.00"


def test_fifty_page_boundary(tmp_path):
    h, p, w = invoice_fixture(50)
    path = tmp_path / "fifty.pdf"
    render_pdf(path, "Fifty page stress test", h, p, w)
    r = extract(path)
    assert r["page_count"] == 50 and len(r["tables"]) == 1
    assert r["row_count"] == 501 and r["headers_removed"] == 49
    assert [row["cells"] for row in r["tables"][0]["rows"]] == sum(p, [])
    assert not r["issues"]


def test_over_limit_rejected(tmp_path):
    doc = fitz.open()
    for _ in range(51):
        doc.new_page()
    path = tmp_path / "over.pdf"
    doc.save(path)
    doc.close()
    with pytest.raises(ValueError, match="50"):
        extract(path)


def test_password_rejected(tmp_path):
    d = fitz.open()
    d.new_page()
    path = tmp_path / "locked.pdf"
    d.save(
        path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret"
    )
    d.close()
    with pytest.raises(ValueError, match="Password"):
        extract(path)


def test_empty_page_requires_review(tmp_path):
    d = fitz.open()
    d.new_page()
    path = tmp_path / "blank.pdf"
    d.save(path)
    d.close()
    r = extract(path)
    assert r["row_count"] == 0 and r["issues"][0]["kind"] == "page_review"


def test_image_only_fixture_uses_real_ocr():
    path = ROOT / "samples" / "02_scanned_inventory.pdf"
    with fitz.open(path) as doc:
        assert all(not p.get_text().strip() for p in doc)
    r = extract(path)
    assert r["page_count"] == 2 and r["row_count"] == 20
    assert len(r["tables"]) == 1 and len(r["tables"][0]["headers"]) == 12
    assert all(p["method"] == "ocr" for p in r["pages"])
    expected = json.loads((ROOT / "samples" / "ground_truth.json").read_text())[
        "02_scanned_inventory.pdf"
    ]["rows"]
    actual = [row["cells"] for row in r["tables"][0]["rows"]]
    # Financial values and identifiers must match exactly, even if text spacing varies.
    assert [[row[c] for c in [0, 1, 4, 5, 6, 7, 8, 9]] for row in actual] == [
        [row[c] for c in [0, 1, 4, 5, 6, 7, 8, 9]] for row in expected
    ]


def test_excel_types_formats_evidence_and_injection():
    r = extract(ROOT / "samples" / "01_invoice_register.pdf")
    r["tables"][0]["rows"][0]["cells"][2] = '=HYPERLINK("https://example.test","bad")'
    wb = load_workbook(io.BytesIO(xlsx(r)))
    assert wb.sheetnames == [
        "Overview",
        "Table 1",
        "Validation",
        "Evidence",
        "Corrections",
    ]
    ws = wb["Table 1"]
    assert ws["A2"].value == "INV-0001" and ws["B2"].value.year == 2026
    assert (
        ws["E2"].value == 2
        and ws["G2"].value == 0.1
        and ws["G2"].number_format == "0.00%"
    )
    assert ws["C2"].data_type == "s" and ws["C2"].value.startswith("=")
    assert ws.freeze_panes == "A2" and len(ws.tables) == 1
    assert wb["Evidence"].max_row == 32
    assert all(c.data_type != "f" for sheet in wb for row in sheet for c in row)


def test_identifier_leading_zeros_preserved():
    assert typed("00123", "Item ID") == "00123"
    assert typed("00123", "Amount") == "00123"
    assert typed("123", "Reference") == "123"
    assert typed("1,234.50", "Amount") == 1234.5
    assert typed("1,234.50", "Invoice amount") == 1234.5


def test_csv_injection_escaped():
    t = {"headers": ["=BAD", "Description"], "rows": [{"cells": ["=1+1", "@payload"]}]}
    data = csv_bytes(t).decode("utf-8-sig")
    assert "'=BAD" in data and "'=1+1" in data and "'@payload" in data


def test_validation_for_empty_misaligned_and_bad_math():
    t = stitch(
        [
            fragment(
                1,
                [
                    ["Item ID", "Quantity", "Unit price", "Net", "Empty"],
                    ["01", "2", "10", "25", ""],
                    ["02", "3", "text", "30", ""],
                    ["03", "", "12", "36", ""],
                ],
            )
        ]
    )
    kinds = {i["kind"] for i in validate(t)}
    assert {"empty_column", "arithmetic", "missing_value"} <= kinds


def test_subtotals_do_not_double_count_grand_total():
    t = stitch(
        [
            fragment(
                1,
                [
                    ["Description", "Amount"],
                    ["First", "10"],
                    ["SUBTOTAL", "10"],
                    ["Second", "20"],
                    ["SUBTOTAL", "20"],
                    ["GRAND TOTAL", "30"],
                ],
            )
        ]
    )
    assert not [i for i in validate(t) if i["kind"] == "total_mismatch"]


def test_line_total_validation():
    t = stitch(
        [
            fragment(
                1,
                [
                    ["Description", "Net", "Tax", "Total"],
                    ["First", "10", "2", "99"],
                    ["Second", "20", "4", "24"],
                ],
            )
        ]
    )
    assert any(i["kind"] == "line_total" and i["row"] == 1 for i in validate(t))


def test_borderless_native_keeps_multiword_headers(tmp_path):
    h, p, w = invoice_fixture(2)
    path = tmp_path / "borderless.pdf"
    render_pdf(path, "Borderless digital invoice", h, p, w, borderless=True)
    r = extract(path)
    assert len(r["tables"]) == 1 and r["tables"][0]["headers"] == h
    assert [row["cells"] for row in r["tables"][0]["rows"]] == sum(p, [])
    assert any(i["kind"] == "alignment_review" for i in r["issues"])
