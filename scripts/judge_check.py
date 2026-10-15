"""Judge smoke checks requiring runtime dependencies only, with honest exit codes."""

import io
import json
import sys
import tempfile
import time
from pathlib import Path

import pymupdf as fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.engine import extract  # noqa: E402 - direct script entry point
from app.exporter import xlsx  # noqa: E402 - direct script entry point


def main():
    started = time.perf_counter()
    checks = []

    def check(name, operation):
        try:
            detail = operation()
            checks.append({"name": name, "passed": True, "detail": detail})
            print("PASS:", name, "-", detail, flush=True)
        except Exception as exc:
            checks.append({"name": name, "passed": False, "detail": str(exc)})
            print("FAIL:", name, "-", exc, flush=True)

    truth = json.loads((ROOT / "samples/ground_truth.json").read_text())
    for filename, expected in truth.items():

        def fixture(filename=filename, expected=expected):
            result = extract(ROOT / "samples" / filename)
            assert len(result["tables"]) == 1, "Unexpected number of tables"
            table = result["tables"][0]
            assert table["headers"] == expected["headers"], "Header mismatch"
            assert [row["cells"] for row in table["rows"]] == expected["rows"], (
                "Data differs from ground truth"
            )
            if "bank" in filename:
                assert {"balance_mismatch", "total_mismatch", "missing_value"} <= {
                    i["kind"] for i in result["issues"]
                }, "Expected source errors were not detected"
            workbook = load_workbook(io.BytesIO(xlsx(result)))
            assert {
                "Overview",
                "Table 1",
                "Validation",
                "Evidence",
                "Corrections",
                "Page coverage",
                "Run log",
            } <= set(workbook.sheetnames)
            assert workbook["Table 1"].max_row == len(expected["rows"]) + 1
            return f"{len(expected['rows'])} rows x {len(expected['headers'])} columns, exact cells/headers; Excel inspected"

        check(filename, fixture)

    complex_truth = json.loads(
        (ROOT / "evaluation/complex/ground_truth.json").read_text(encoding="utf-8")
    )
    for name in ("nested_line_items", "row_split_across_pages"):

        def complex_fixture(name=name):
            result = extract(ROOT / "evaluation/complex/pdfs" / f"{name}.pdf")
            actual = [
                {"headers": t["headers"], "rows": [r["cells"] for r in t["rows"]]}
                for t in result["tables"]
            ]
            assert actual == complex_truth[name]["tables"], "Complex table mismatch"
            workbook = load_workbook(io.BytesIO(xlsx(result)))
            assert "Structure" in workbook.sheetnames
            if name == "nested_line_items":
                assert all(
                    r.get("inherited_cells") for r in result["tables"][0]["rows"]
                )
                return "Nested child headers and three shared order IDs preserved; Structure sheet inspected"
            row = result["tables"][0]["rows"][2]
            assert [part["page"] for part in row["sources"]] == [1, 2]
            assert workbook["Evidence"].max_row == 8, "Missing source part in Excel"
            return "One logical continuation row retains two source-page parts; six rows x 12 columns"

        check(name.replace("_", " "), complex_fixture)

    (ROOT / "tmp").mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="damnlens-judge-", dir=ROOT / "tmp"
    ) as temp:
        temp = Path(temp)

        def fifty():
            document = fitz.open()
            with fitz.open(ROOT / "samples/01_invoice_register.pdf") as source:
                for _ in range(50):
                    document.insert_pdf(source, from_page=0, to_page=0)
            path = temp / "fifty.pdf"
            document.save(path)
            document.close()
            result = extract(path)
            expected = truth["01_invoice_register.pdf"]["rows"][:10] * 50
            assert len(result["tables"]) == 1
            assert [r["cells"] for r in result["tables"][0]["rows"]] == expected
            assert result["headers_removed"] == 49
            return "50 pages, 500 rows, 12 columns; 49 headers removed; duplicate business rows preserved"

        check("50-page boundary and duplicate preservation", fifty)

        def rejected():
            document = fitz.open()
            for _ in range(51):
                document.new_page()
            path = temp / "over.pdf"
            document.save(path)
            document.close()
            try:
                extract(path)
            except ValueError as exc:
                assert "50" in str(exc)
                return "51-page input rejected with an actionable message"
            raise AssertionError("An oversized document was accepted")

        check("Over-limit recovery", rejected)

        def resumed():
            source = ROOT / "samples/01_invoice_register.pdf"
            cache = temp / "checkpoints"

            def interrupt(done, total, message):
                if done == 2 and message == "Inspecting page 3":
                    raise RuntimeError("Deliberate judge-check interruption")

            try:
                extract(source, checkpoint_dir=cache, progress=interrupt)
            except RuntimeError as exc:
                assert "Deliberate judge-check" in str(exc)
            else:
                raise AssertionError("Interruption was not exercised")
            result = extract(source, checkpoint_dir=cache, resume=True)
            assert [p.get("resumed", False) for p in result["pages"]] == [
                True,
                True,
                False,
            ]
            assert [r["cells"] for r in result["tables"][0]["rows"]] == truth[
                "01_invoice_register.pdf"
            ]["rows"]
            return "Two verified pages reused after interruption; final page extracted; exact rows retained"

        check("Verified page checkpoint recovery", resumed)

    report = {
        "application": "DamnLens AI Agent",
        "seconds": round(time.perf_counter() - started, 2),
        "scope": "Bundled synthetic samples, complex structure, source parts, XLSX readability, page boundaries and checkpoint recovery. This is a reproducible smoke check, not a hidden-PDF accuracy guarantee.",
        "checks": checks,
    }
    (ROOT / "reports").mkdir(exist_ok=True)
    (ROOT / "reports/judge-check-latest.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8"
    )
    print(
        f"\n{sum(c['passed'] for c in checks)}/{len(checks)} checks passed in {report['seconds']}s. Report: reports/judge-check-latest.json"
    )
    return 0 if all(c["passed"] for c in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
