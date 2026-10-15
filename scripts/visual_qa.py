"""Render source fixtures and inspect exported workbook structure."""

from pathlib import Path
import io
import pymupdf as fitz
from PIL import Image, ImageDraw
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
out = ROOT / "tmp" / "visual"
out.mkdir(parents=True, exist_ok=True)
sheet = Image.new("RGB", (1120, 3 * 740), "#f8f9f5")
draw = ImageDraw.Draw(sheet)
for i, path in enumerate(sorted((ROOT / "samples").glob("*.pdf"))):
    with fitz.open(path) as doc:
        image = Image.open(
            io.BytesIO(doc[0].get_pixmap(dpi=72, alpha=False).tobytes("png"))
        ).convert("RGB")
        sheet.paste(image, (0, i * 740 + 30))
        draw.text((20, i * 740 + 8), path.name, fill="#173c37")
    x = ROOT / "samples" / "converted" / path.with_suffix(".xlsx").name
    wb = load_workbook(x)
    print(
        path.name,
        wb.sheetnames,
        wb["Table 1"].max_row - 1,
        "rows",
        wb["Table 1"].max_column - 4,
        "data columns",
    )
sheet.save(out / "sample-contact-sheet.png")
print("Rendered", out / "sample-contact-sheet.png")
